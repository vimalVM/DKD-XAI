"""
Script to generate a professional, beautifully formatted Word document (.docx)
summarizing the DKD-XAI project from start to finish, including Phase 1 & Phase 2 outcomes.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for table cells."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Set light borders for standard tables."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout(doc, text_list, title="KEY TAKEAWAY", border_color="0D9488", bg_color="F0FDFA"):
    """Create an elegant highlighted callout block."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run_title = p.add_run(f"📌 {title}: ")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(13, 148, 136) if border_color == "0D9488" else RGBColor(27, 54, 93)
    
    for i, line in enumerate(text_list):
        if i == 0:
            run_text = p.add_run(line)
            run_text.font.name = "Arial"
            run_text.font.size = Pt(9.5)
            run_text.font.color.rgb = RGBColor(30, 41, 59)
        else:
            p_sub = cell.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(2)
            p_sub.paragraph_format.space_after = Pt(2)
            p_sub.paragraph_format.line_spacing = 1.15
            run_text = p_sub.add_run(line)
            run_text.font.name = "Arial"
            run_text.font.size = Pt(9.5)
            run_text.font.color.rgb = RGBColor(30, 41, 59)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_document():
    doc = docx.Document()
    
    # Page setup: Standard Letter, 0.8 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    
    # Colors
    NAVY = RGBColor(27, 54, 93)      # #1B365D - Primary headers
    TEAL = RGBColor(13, 148, 136)    # #0D9488 - Sub-headers / accents
    CHARCOAL = RGBColor(51, 65, 85)  # #334155 - Body text
    MUTED = RGBColor(100, 116, 139)  # #64748B - Metadata
    
    # Title Block
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run("DKD-XAI: Explainable AI for Early Diabetic Kidney Disease Risk Stratification")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.bold = True
    run_title.font.color.rgb = NAVY
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(8)
    run_sub = p_sub.add_run("Project Summary & Outcomes Report | End-to-End Overview")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = TEAL
    run_sub.bold = True
    
    # Metadata bar
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(14)
    run_meta = p_meta.add_run("Author: Vimal Vijayan (Roll No. 17030925059)  •  Mentor: Dr. Mayura Nagar  •  Date: October 2026")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(9.5)
    run_meta.font.color.rgb = MUTED
    
    # Horizontal divider rule
    p_divider = doc.add_paragraph()
    p_divider.paragraph_format.space_after = Pt(8)
    p_divider_run = p_divider.add_run("―" * 58)
    p_divider_run.font.color.rgb = RGBColor(226, 232, 240)
    p_divider_run.bold = True

    # ----------------------------------------------------
    # SECTION 1: EXECUTIVE SUMMARY
    # ----------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("1. Executive Summary")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = NAVY
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(
        "Diabetic Kidney Disease (DKD) is one of the most severe microvascular complications of diabetes, "
        "affecting nearly 40% of diabetic patients and standing as the leading cause of end-stage renal disease globally. "
        "Traditionally, DKD is diagnosed only after kidney damage has already occurred, using laboratory markers like "
        "elevated Urine Albumin-to-Creatinine Ratio (UACR ≥ 30 mg/g) or reduced estimated Glomerular Filtration Rate (eGFR < 60 mL/min/1.73m²)."
    )
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL
    
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(
        "The DKD-XAI project provides an end-to-end clinical machine learning and generative explainable AI (XAI) solution. "
        "Its central breakthrough is predicting kidney damage risk early using exclusively non-kidney routine parameters "
        "(e.g., blood pressure, age, HbA1c, lipids, diabetes duration) while strictly eliminating data leakage. "
        "The project was executed in two systematic phases:"
    )
    r2.font.name = "Arial"
    r2.font.size = Pt(10)
    r2.font.color.rgb = CHARCOAL

    bullet1 = doc.add_paragraph(style='List Bullet')
    bullet1.paragraph_format.space_after = Pt(3)
    bullet1.paragraph_format.line_spacing = 1.15
    r = bullet1.add_run("Phase 1 (Predictive Pipeline & Explainability): ")
    r.bold = True
    r.font.color.rgb = NAVY
    r = bullet1.add_run("Harmonized 4,024 adult diabetic patients across 4 NHANES cycles, benchmarked four ML/DL algorithms "
                        "(Logistic Regression, Random Forest, XGBoost, and PyTorch Deep Learning MLP), established chronological temporal validation, "
                        "and calculated transparent SHAP feature attributions.")
    r.font.color.rgb = CHARCOAL

    bullet2 = doc.add_paragraph(style='List Bullet')
    bullet2.paragraph_format.space_after = Pt(8)
    bullet2.paragraph_format.line_spacing = 1.15
    r = bullet2.add_run("Phase 2 (Evidence-Grounded AI Agent & Clinical Web Interface): ")
    r.bold = True
    r.font.color.rgb = NAVY
    r = bullet2.add_run("Engineered a Groq-powered clinical reasoning agent (Llama 3.3 70B) grounded in official ADA 2026 / KDIGO 2022 guidelines "
                        "with strict automated faithfulness checks, paired with a full-stack interactive web application featuring natural-language vitals extraction.")
    r.font.color.rgb = CHARCOAL

    add_callout(
        doc,
        ["• Best Overall Model: XGBoost achieved 0.732 AUROC on holdout testing and 0.723 on unseen future temporal cycles.",
         "• Strict Zero Leakage: Serum creatinine, eGFR, and UACR are strictly forbidden as model inputs.",
         "• Trustworthy AI: Generative explanations cite verifiable medical guidelines with zero statistical hallucinations."],
        title="CORE HIGHLIGHTS"
    )

    # ----------------------------------------------------
    # SECTION 2: DATASET & PREPROCESSING
    # ----------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("2. Cohort Data & Zero-Leakage Architecture")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(
        "Data was pooled from the U.S. Centers for Disease Control (CDC) National Health and Nutrition Examination Survey (NHANES), "
        "spanning four cycles from 2013 to 2023. Only adult participants (18+) with verified diabetes (doctor-diagnosed or HbA1c ≥ 6.5%) "
        "who had complete laboratory measurements were included."
    )
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL

    # Table of Dataset Cohort
    table_cohort = doc.add_table(rows=6, cols=3)
    table_cohort.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_cohort.autofit = False
    set_table_borders(table_cohort)
    
    col_widths = [Inches(2.5), Inches(1.8), Inches(2.2)]
    headers = ["NHANES Survey Cycle", "Participant Count", "Status / Pipeline Usage"]
    for i, h in enumerate(headers):
        cell = table_cohort.cell(0, i)
        cell.width = col_widths[i]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    cohort_data = [
        ("2013–2014 (_H cycle)", "790 adults", "Reconstructed from raw XPT files"),
        ("2015–2016 (_I cycle)", "892 adults", "Reconstructed from raw XPT files"),
        ("2017–March 2020 (P_ cycle)", "1,455 adults", "Standardized baseline cohort"),
        ("Aug 2021–Aug 2023 (_L cycle)", "887 adults", "Temporal holdout evaluation set"),
        ("Total Unified Cohort", "4,024 adults", "Preserved 100% real records (no synthesis)")
    ]

    for row_idx, data in enumerate(cohort_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        if row_idx == 5:
            bg = "E2E8F0"
        for col_idx, text in enumerate(data):
            cell = table_cohort.cell(row_idx, col_idx)
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.size = Pt(9)
            if row_idx == 5:
                run.bold = True
                run.font.color.rgb = NAVY
            else:
                run.font.color.rgb = CHARCOAL

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Features and Labels
    p_feat = doc.add_paragraph()
    p_feat.paragraph_format.space_after = Pt(4)
    p_feat.paragraph_format.line_spacing = 1.15
    r = p_feat.add_run("The 13 Non-Kidney Input Features: ").bold = True
    r = p_feat.add_run("Age, Biological Sex (Female), HbA1c (%), Body Mass Index (BMI), Waist Circumference (cm), "
                       "Systolic Blood Pressure (SBP), Diastolic Blood Pressure (DBP), HDL Cholesterol, Total Cholesterol, "
                       "Diagnosed Diabetes Indicator, Insulin Usage, Oral Diabetes Pill Usage, and Diabetes Duration in Years.")

    p_targets = doc.add_paragraph()
    p_targets.paragraph_format.space_after = Pt(6)
    p_targets.paragraph_format.line_spacing = 1.15
    r = p_targets.add_run("The Two Clinical Target Tasks: ").bold = True
    r = p_targets.add_run(
        "\n1. Any Kidney Damage (dkd_any): Binary prediction (1,578 positive vs 2,446 negative; 39.2% prevalence) where "
        "positive means single-visit UACR ≥ 30 mg/g OR eGFR < 60 mL/min/1.73m²."
        "\n2. Early DKD vs None (early_dkd): Distinguishes early microalbuminuria with preserved filtration (UACR 30–300 mg/g and eGFR ≥ 60) "
        "from healthy kidneys (699 early cases vs 2,446 controls; 22.2% positive among 3,145 eligible rows; 879 advanced cases excluded)."
    )

    # ----------------------------------------------------
    # SECTION 3: PHASE 1 METHODOLOGY & OUTCOMES
    # ----------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("3. Phase 1: Machine Learning & Deep Learning Pipeline")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(
        "Phase 1 established an empirical comparison across four diverse model architectures: "
        "Linear (Logistic Regression with L2 regularization), Tree Ensemble (Random Forest and XGBoost), and Deep Learning (PyTorch Multi-Layer Perceptron). "
        "The pipeline incorporated 80/20 stratified holdout splits, 5-fold cross-validation, Youden's J threshold optimization on out-of-fold training data, "
        "and SHAP (SHapley Additive exPlanations) explainability."
    )
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL

    # Outcomes Table: dkd_any
    h2_1 = doc.add_paragraph()
    h2_1.paragraph_format.space_before = Pt(6)
    h2_1.paragraph_format.space_after = Pt(3)
    r = h2_1.add_run("Outcome A: Any Kidney Damage (dkd_any) — Holdout Test Performance (N = 805)")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = TEAL

    table_any = doc.add_table(rows=5, cols=8)
    table_any.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_any.autofit = False
    set_table_borders(table_any)

    cols_w_any = [Inches(1.5), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.7), Inches(0.7), Inches(0.6), Inches(0.6)]
    headers_any = ["Model Architecture", "CV AUROC", "Test AUROC", "95% CI", "Sensitivity", "Specificity", "Accuracy", "Threshold"]
    for i, h in enumerate(headers_any):
        cell = table_any.cell(0, i)
        cell.width = cols_w_any[i]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    any_results = [
        ("XGBoost (Selected Best)", "0.732", "0.732", "0.699–0.768", "62.7%", "72.2%", "68.4%", "0.513"),
        ("Random Forest", "0.729", "0.727", "0.693–0.762", "59.8%", "72.8%", "67.7%", "0.501"),
        ("PyTorch MLP (DL)", "0.723", "0.729", "0.694–0.762", "70.9%", "63.6%", "66.5%", "0.474"),
        ("Logistic Regression", "0.719", "0.722", "0.687–0.757", "60.4%", "72.8%", "68.0%", "0.527")
    ]

    for row_idx, row in enumerate(any_results, start=1):
        bg = "F0FDFA" if row_idx == 1 else ("F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        for col_idx, val in enumerate(row):
            cell = table_any.cell(row_idx, col_idx)
            cell.width = cols_w_any[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            if row_idx == 1:
                run.bold = True
                run.font.color.rgb = RGBColor(15, 118, 110)
            else:
                run.font.color.rgb = CHARCOAL

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Outcomes Table: early_dkd
    h2_2 = doc.add_paragraph()
    h2_2.paragraph_format.space_before = Pt(6)
    h2_2.paragraph_format.space_after = Pt(3)
    r = h2_2.add_run("Outcome B: Early DKD vs None (early_dkd) — Holdout Test Performance (N = 629)")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = TEAL

    table_early = doc.add_table(rows=5, cols=8)
    table_early.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_early.autofit = False
    set_table_borders(table_early)

    for i, h in enumerate(headers_any):
        cell = table_early.cell(0, i)
        cell.width = cols_w_any[i]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    early_results = [
        ("PyTorch MLP (Selected Best)", "0.678", "0.643", "0.590–0.693", "52.9%", "69.1%", "65.5%", "0.507"),
        ("Random Forest", "0.678", "0.657", "0.601–0.707", "48.6%", "70.1%", "65.3%", "0.510"),
        ("Logistic Regression", "0.676", "0.644", "0.589–0.696", "50.7%", "68.3%", "64.4%", "0.506"),
        ("XGBoost", "0.674", "0.666", "0.612–0.713", "62.9%", "62.2%", "62.3%", "0.486")
    ]

    for row_idx, row in enumerate(early_results, start=1):
        bg = "F0FDFA" if row_idx == 1 else ("F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        for col_idx, val in enumerate(row):
            cell = table_early.cell(row_idx, col_idx)
            cell.width = cols_w_any[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            if row_idx == 1:
                run.bold = True
                run.font.color.rgb = RGBColor(15, 118, 110)
            else:
                run.font.color.rgb = CHARCOAL

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Chronological Temporal Validation
    h2_3 = doc.add_paragraph()
    h2_3.paragraph_format.space_before = Pt(6)
    h2_3.paragraph_format.space_after = Pt(3)
    r = h2_3.add_run("Chronological Temporal Validation (Trained on 2013–2020, Tested on 2021–2023)")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = TEAL

    p_temp = doc.add_paragraph()
    p_temp.paragraph_format.space_after = Pt(6)
    p_temp.paragraph_format.line_spacing = 1.15
    r = p_temp.add_run(
        "To rigorously test model generalization against cohort drift over time, all models were tuned and trained exclusively "
        "on the historical cycles (2013–2020, N = 3,137) and evaluated on the subsequent cycle (2021–2023, N = 887). "
        "XGBoost maintained a strong temporal AUROC of 0.7230 (AUPRC: 0.6653, Sensitivity: 62.7%, Specificity: 70.8%), "
        "confirming that the learned risk patterns are durable and transportable across real-world survey cohorts."
    )
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = CHARCOAL

    # Explainability & Key Drivers
    p_shap = doc.add_paragraph()
    p_shap.paragraph_format.space_after = Pt(6)
    p_shap.paragraph_format.line_spacing = 1.15
    r = p_shap.add_run("Top Clinical Risk Drivers Revealed by SHAP: ").bold = True
    r = p_shap.add_run(
        "Across all tree and linear attributions, the dominant contributors to elevated kidney damage risk were: "
        "(1) Older Age, (2) Elevated Glycated Hemoglobin (HbA1c > 8.0%), (3) High Systolic Blood Pressure (SBP > 140 mmHg), "
        "(4) Longer Diabetes Duration (> 10 years), and (5) High Body Mass Index (BMI > 30 kg/m²). "
        "Protective features lowering predicted risk included well-managed SBP (< 120 mmHg) and normal HbA1c (< 6.5%)."
    )
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = CHARCOAL

    # Embed SHAP and ROC plots if available
    roc_img_path = r"d:\DKD-XAI_Updated_Full\DKD-XAI_Updated\results\dkd_any\roc_pr_calibration.png"
    shap_img_path = r"d:\DKD-XAI_Updated_Full\DKD-XAI_Updated\results\dkd_any\shap_summary_beeswarm.png"
    
    if os.path.exists(shap_img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        doc.add_picture(shap_img_path, width=Inches(5.6))
        p_caption = doc.add_paragraph()
        p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_caption.paragraph_format.space_after = Pt(8)
        rc = p_caption.add_run("Figure 1: SHAP Beeswarm Plot — Feature attributions and directionality for XGBoost on the test cohort.")
        rc.font.size = Pt(8.5)
        rc.font.italic = True
        rc.font.color.rgb = MUTED

    # ----------------------------------------------------
    # SECTION 4: PHASE 2 METHODOLOGY & OUTCOMES
    # ----------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("4. Phase 2: Evidence-Grounded AI Agent & Interactive Web Platform")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(
        "A critical failure mode of clinical AI is the 'black-box' dilemma: clinicians cannot act on unexplained risk numbers, "
        "and generic generative LLMs tend to hallucinate inaccurate medical advice or make ungrounded statistical claims. "
        "Phase 2 tackled this directly by building a specialized, multi-tiered Explainable AI Clinical System."
    )
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL

    bullet_agent1 = doc.add_paragraph(style='List Bullet')
    bullet_agent1.paragraph_format.space_after = Pt(3)
    bullet_agent1.paragraph_format.line_spacing = 1.15
    r = bullet_agent1.add_run("Dynamic Model Binding: ")
    r.bold = True
    r.font.color.rgb = NAVY
    r = bullet_agent1.add_run("The agent dynamically loads whichever model was selected by Phase 1 (XGBoost or PyTorch MLP), "
                              "calculates the exact individual risk probability, and extracts patient-specific positive and negative SHAP factors.")
    r.font.color.rgb = CHARCOAL

    bullet_agent2 = doc.add_paragraph(style='List Bullet')
    bullet_agent2.paragraph_format.space_after = Pt(3)
    bullet_agent2.paragraph_format.line_spacing = 1.15
    r = bullet_agent2.add_run("Evidence Grounding (dkd_evidence_base.json): ")
    r.bold = True
    r.font.color.rgb = NAVY
    r = bullet_agent2.add_run("For each top risk driver, the agent retrieves verified clinical guidelines from the American Diabetes Association "
                              "(ADA Standards of Care 2026) and KDIGO 2022 Guidelines. The LLM is restricted to citing only these verified sources.")
    r.font.color.rgb = CHARCOAL

    bullet_agent3 = doc.add_paragraph(style='List Bullet')
    bullet_agent3.paragraph_format.space_after = Pt(3)
    bullet_agent3.paragraph_format.line_spacing = 1.15
    r = bullet_agent3.add_run("Groq-Powered LLM Reasoning (Llama 3.3 70B): ")
    r.bold = True
    r.font.color.rgb = NAVY
    r = bullet_agent3.add_run("Generates concise (~120 words), plain-English clinical risk summaries explaining why the patient is at risk, "
                              "attaching mandatory source citations, and ending with a clear reminder that this is a risk estimate, not a diagnosis.")
    r.font.color.rgb = CHARCOAL

    bullet_agent4 = doc.add_paragraph(style='List Bullet')
    bullet_agent4.paragraph_format.space_after = Pt(3)
    bullet_agent4.paragraph_format.line_spacing = 1.15
    r = bullet_agent4.add_run("Automated Faithfulness Verification: ")
    r.bold = True
    r.font.color.rgb = NAVY
    r = bullet_agent4.add_run("Every LLM output undergoes a deterministic post-generation check to confirm: (1) exact risk % is preserved, "
                              "(2) no ungrounded numbers were introduced, and (3) all citations correspond exactly to retrieved sources.")
    r.font.color.rgb = CHARCOAL

    bullet_agent5 = doc.add_paragraph(style='List Bullet')
    bullet_agent5.paragraph_format.space_after = Pt(6)
    bullet_agent5.paragraph_format.line_spacing = 1.15
    r = bullet_agent5.add_run("Deterministic Offline Fallback: ")
    r.bold = True
    r.font.color.rgb = NAVY
    r = bullet_agent5.add_run("If the Groq API is offline or without internet access, a rule-based template generates identical, "
                              "100% faithful explanations with zero downtime.")
    r.font.color.rgb = CHARCOAL

    # Full Stack Web Application
    h2_web = doc.add_paragraph()
    h2_web.paragraph_format.space_before = Pt(6)
    h2_web.paragraph_format.space_after = Pt(3)
    r = h2_web.add_run("Phase 2 Product: Interactive Clinical Web Application")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = TEAL

    p_web = doc.add_paragraph()
    p_web.paragraph_format.space_after = Pt(6)
    p_web.paragraph_format.line_spacing = 1.15
    r = p_web.add_run(
        "To make the predictive engine accessible to clinicians, a production-grade web dashboard was developed:\n"
        "• FastAPI Backend (backend_server.py): Provides REST endpoints for live model inference, natural language vitals extraction, "
        "and AI agent explanation generation.\n"
        "• React/Vite Frontend (frontend/): Built with healthcare-grade UI (soft slate, medical teal #0D9488, clinical typography).\n"
        "• Dual Input Modes: Traditional structured clinical vitals form alongside a Conversational Natural Language Assistant where clinicians "
        "can paste unstructured clinical shorthand (e.g., '65M, T2D for 15y, A1c 9.2, BP 155/90, BMI 31') which is automatically parsed into features.\n"
        "• Visual SHAP Attribution Meters: Dynamic color-coded bar meters displaying exact risk-increasing factors in coral-red and protective factors in emerald-green.\n"
        "• Pre-loaded Benchmark Patients: 4 real clinical test cases (e.g., High-Risk Longstanding Diabetic, Well-Controlled Early Patient) selectable via a 1-click header dropdown."
    )
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = CHARCOAL

    add_callout(
        doc,
        ["• Faithfulness Checker: Catches altered percentages, invented clinical facts, and ungrounded citations.",
         "• Clinician Workflow: Translates complex tabular data and model math into an actionable 1-minute risk narrative.",
         "• Verified Sources: Every claim links to official ADA and KDIGO clinical guidelines."],
        title="PHASE 2 OUTCOMES & SAFETY"
    )

    # ----------------------------------------------------
    # SECTION 5: CLINICAL LIMITATIONS & BEST PRACTICES
    # ----------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("5. Scientific Limitations & Ethical Cautions")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = NAVY

    p_lim = doc.add_paragraph()
    p_lim.paragraph_format.space_after = Pt(6)
    p_lim.paragraph_format.line_spacing = 1.15
    r = p_lim.add_run(
        "For complete scientific transparency and clinical safety, the following boundaries are formally defined in the project:"
    )
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL

    lim_items = [
        ("Cross-Sectional Marker Proxies: ", "NHANES measures UACR and eGFR at a single survey visit. Clinical diagnostic criteria for CKD require persistent abnormalities for > 3 months. Therefore, labels are single-visit risk markers, not confirmed diagnoses."),
        ("No National Prevalence Claims: ", "The pipeline evaluates predictive ML algorithms without applying NHANES complex survey weights, strata, or clusters. Metrics reflect cohort performance, not US population prevalence."),
        ("SHAP is Attribution, Not Causation: ", "SHAP values describe how the model weighted a given feature for its statistical score. They do not demonstrate that altering that feature (e.g. lowering waist circumference) will directly cure or halt kidney disease."),
        ("Research Prototype Status: ", "This system is an academic research prototype designed to support clinical triage and risk awareness; it is not a standalone diagnostic medical device.")
    ]

    for title, desc in lim_items:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        bp.paragraph_format.line_spacing = 1.15
        r = bp.add_run(title)
        r.bold = True
        r.font.color.rgb = NAVY
        r = bp.add_run(desc)
        r.font.color.rgb = CHARCOAL

    # ----------------------------------------------------
    # SECTION 6: SUMMARY COMPARISON TABLE
    # ----------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)
    r = h1.add_run("6. Comprehensive Phase 1 vs Phase 2 Summary")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = NAVY

    table_comp = doc.add_table(rows=8, cols=3)
    table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_comp.autofit = False
    set_table_borders(table_comp)

    col_w_comp = [Inches(1.8), Inches(2.4), Inches(2.4)]
    comp_headers = ["Project Dimension", "Phase 1: Machine Learning & XAI", "Phase 2: AI Agent & Web App"]
    for i, h in enumerate(comp_headers):
        cell = table_comp.cell(0, i)
        cell.width = col_w_comp[i]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    comp_rows = [
        ("Core Objective", "Predict kidney damage using routine vitals; compare ML vs DL algorithms.", "Translate predictions into faithful, guideline-backed explanations and UI."),
        ("Models / Tools", "Logistic Regression, Random Forest, XGBoost, PyTorch MLP, TreeSHAP.", "Groq API (Llama 3.3 70B), FastAPI, React/Vite, Faithfulness Checker."),
        ("Key Evaluation", "Holdout CV, 95% Bootstrap CI, Chronological Temporal Holdout (2021–2023).", "Citation verification, exact risk % consistency, latency, usability."),
        ("Best Model Selected", "XGBoost (dkd_any: 0.732 AUROC); PyTorch MLP (early_dkd: 0.678 CV AUROC).", "Loads whichever model Phase 1 selected dynamically via joblib metadata."),
        ("Key Output / Artifact", "Trained models (.joblib), ROC/PR curves, SHAP beeswarm & dependence plots.", "dkd_ai_agent.py, backend_server.py, interactive clinical web dashboard."),
        ("Verification Guardrail", "Stratified CV, leakage test (validate_project.py), Youden's threshold.", "Faithfulness checker (flags hallucinations, uncited claims, altered numbers)."),
        ("Final Clinical Deliverable", "Validated mathematical risk score and top contributing numerical factors.", "Interactive patient narrative report with guideline citations and interactive UI.")
    ]

    for row_idx, data in enumerate(comp_rows, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(data):
            cell = table_comp.cell(row_idx, col_idx)
            cell.width = col_w_comp[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            if col_idx == 0:
                run.bold = True
                run.font.color.rgb = NAVY
            else:
                run.font.color.rgb = CHARCOAL

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Concluding signature / sign-off
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_sign.add_run("― End of DKD-XAI Comprehensive Project Summary ―\n")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = MUTED
    r.italic = True
    
    r_sub = p_sign.add_run("Prepared for Academic & Clinical Review  •  October 2026")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(8.5)
    r_sub.font.color.rgb = MUTED

    output_path = r"d:\DKD-XAI_Updated_Full\DKD-XAI_Updated\DKD_XAI_Complete_Project_Summary.docx"
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    build_document()
